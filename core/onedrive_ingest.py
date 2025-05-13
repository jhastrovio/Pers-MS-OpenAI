import os
import asyncio
from core.graph_client import MSGraphClient
from core.auth import MSGraphAuth
from core.vectorstore_upload import upload_to_vectorstore
import httpx
from pathlib import Path
from datetime import datetime
import re
import PyPDF2
from docx import Document as DocxDocument
from openpyxl import load_workbook
from pptx import Presentation as PptxPresentation

# Ingest all files from the PERS-MS-OPENAI root folder and subfolders
TARGET_FOLDERS = [
    "PERS-MS-OPENAI",
]

DOWNLOAD_DIR = Path("downloads_onedrive")
DOWNLOAD_DIR.mkdir(exist_ok=True)

# Set your vector store ID here (reuse this for all uploads)
VECTOR_STORE_ID = os.getenv("OPENAI_VECTOR_STORE_ID") or "vs_abc123..."  # TODO: Replace with your actual vector store ID

def extract_internal_author(file_path):
    ext = os.path.splitext(file_path)[1].lower()
    try:
        if ext == ".docx":
            doc = DocxDocument(file_path)
            props = doc.core_properties
            return props.author or ""
        elif ext == ".xlsx":
            wb = load_workbook(file_path, read_only=True)
            props = wb.properties
            return props.creator or ""
        elif ext == ".pptx":
            prs = PptxPresentation(file_path)
            props = prs.core_properties
            return props.author or ""
        elif ext == ".pdf":
            with open(file_path, "rb") as f:
                reader = PyPDF2.PdfReader(f)
                info = reader.metadata
                return info.get("/Author", "") if info else ""
    except Exception as e:
        print(f"[WARN] Could not extract internal author from {file_path}: {e}")
    return ""

async def main():
    user_email = os.environ.get("USER_EMAIL")
    if not user_email:
        raise ValueError("USER_EMAIL environment variable not set.")
    auth = MSGraphAuth(
        client_id=os.environ["CLIENT_ID"],
        client_secret=os.environ["CLIENT_SECRET"],
        tenant_id=os.environ["TENANT_ID"]
    )
    client = MSGraphClient(auth=auth)
    for folder in TARGET_FOLDERS:
        print(f"[INFO] Ingesting OneDrive folder: {folder}")
        items = await client.fetch_drive_items(user_email, folder_path=folder, recursive=True)
        for item in items:
            if "file" in item:
                print(f"[FILE] {item['full_path']}")
                # Download file locally
                file_id = item["id"]
                filename = item["name"]
                local_path = DOWNLOAD_DIR / filename
                # Download file content from OneDrive
                try:
                    content = await client.get_file_content(file_id, user_email=user_email)
                    # Save as binary or text depending on type
                    mode = "wb" if isinstance(content, bytes) else "w"
                    with open(local_path, mode, encoding=None if mode=="wb" else "utf-8") as f:
                        f.write(content)
                except Exception as e:
                    print(f"[WARN] Failed to download {filename} from OneDrive: {e}")
                    if hasattr(e, 'response') and hasattr(e.response, 'text'):
                        print(f"[DEBUG] Error response: {e.response.text}")
                    continue
                # After download, try to extract internal author
                internal_author = extract_internal_author(local_path)
                if internal_author:
                    author = internal_author
                else:
                    author = item.get("createdBy", {}).get("user", {}).get("displayName", "")
                # Upload to OpenAI Vector Store
                match = re.match(r"([A-Za-z0-9_-]+)_(.+)", filename)
                if match:
                    email_id = match.group(1)
                    original_filename = match.group(2)
                else:
                    email_id = None
                    original_filename = filename

                metadata = {
                    "id": item["id"],
                    "name": original_filename,
                    "type": "file",
                    "extension": os.path.splitext(filename)[1].lower(),
                    "size": item.get("size"),
                    "author": author,
                    "last_modified_by": item.get("lastModifiedBy", {}).get("user", {}).get("displayName", ""),
                    "date": item.get("lastModifiedDateTime"),
                    "path": item.get("full_path"),
                    "url": item.get("webUrl", "")
                }
                try:
                    openai_file_id, _ = upload_to_vectorstore(
                        file_path=str(local_path),
                        attributes=metadata,
                        vector_store_id=VECTOR_STORE_ID
                    )
                    print(f"[UPLOAD] {filename} → OpenAI file_id: {openai_file_id}")
                except Exception as e:
                    print(f"[WARN] Failed to upload {filename} to OpenAI Vector Store: {e}")

if __name__ == "__main__":
    asyncio.run(main()) 